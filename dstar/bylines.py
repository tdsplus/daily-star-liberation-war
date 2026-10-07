"""Real writers for articles the site bylines only as "The Daily Star".

Many older pieces (and most reprints) carry the generic byline "The Daily Star"
in their metadata, while the page itself names the writer. The name is usually
in one of these places:
  * the standfirst printed under the headline ("...argues Kajalie Shehreen Islam")
  * a byline line at the top of the text
  * a bio line at the end ("Mofidul Hoque is Trustee, Liberation War Museum.")
  * the writer's e-mail address at the end

The e-mail addresses are hidden by Cloudflare's e-mail obfuscation, which every
browser decodes when it shows the page. `emails()` does the same decoding on the
pages already in the cache (no new requests), only so that the writer's name can
be read from the address. Addresses are not written to the workbook.

Decisions are made by hand from this evidence and kept in
discovery/author_fixes.csv (status: named / editorial / not named). A name is
used only when the page states it; a role on its own ("The writer is Executive
Editor") is noted, not turned into a name.

Usage:
  python -m dstar.bylines pending            # generic-byline rows with no decision yet
  python -m dstar.bylines evidence <id> ...  # clues for given node ids
  python -m dstar.bylines apply <in.xlsx> <out.xlsx>   # fix an exported workbook
"""
import csv
import json
import os
import re
import sys

from bs4 import BeautifulSoup

from .extract import load_records, text_path
from .store import ROOT, node_id

FIXES = os.path.join(ROOT, "discovery", "author_fixes.csv")
FETCH_LOG = os.path.join(ROOT, "cache", "fetch_log.jsonl")
GENERIC = {"the daily star"}
NAMED, EDITORIAL, NOT_NAMED = "named", "editorial", "not named"


def is_generic(author):
    return (author or "").strip().lower() in GENERIC


def load_fixes():
    if not os.path.exists(FIXES):
        return {}
    with open(FIXES, newline="", encoding="utf-8") as fh:
        return {r["node_id"]: r for r in csv.DictReader(fh)}


def note(fix):
    """Text for the 'Author note' column: where the name was found, or why there is none."""
    return fix["evidence"]


# --- evidence ---------------------------------------------------------------

_CACHE = None


def _html(rec):
    global _CACHE
    if _CACHE is None:
        _CACHE = {}
        if os.path.exists(FETCH_LOG):
            with open(FETCH_LOG, encoding="utf-8") as fh:
                for line in fh:
                    if line.strip():
                        e = json.loads(line)
                        if e.get("cache"):
                            _CACHE[e["url"]] = _CACHE[e.get("final_url") or e["url"]] = e["cache"]
    for u in (rec["fetched_url"], rec["canonical_url"]):
        path = os.path.join(ROOT, _CACHE.get(u, "-"))
        if os.path.exists(path):
            with open(path, encoding="utf-8", errors="replace") as fh:
                return fh.read()
    return ""


def _cf_decode(hexs):
    """Cloudflare e-mail obfuscation: first byte is an XOR key for the rest."""
    key = int(hexs[:2], 16)
    return "".join(chr(int(hexs[i:i + 2], 16) ^ key) for i in range(2, len(hexs), 2))


def emails(soup):
    found = [_cf_decode(el["data-cfemail"]) for el in soup.select("[data-cfemail]")]
    found += [_cf_decode(a["href"].split("#", 1)[1])
              for a in soup.select('a[href*="/cdn-cgi/l/email-protection#"]')]
    return list(dict.fromkeys(found))


ROLES = (r"writer|journalist|professor|lecturer|columnist|editor|freedom fighter|researcher|analyst|"
         r"author|trustee|member|teacher|student|former|retired|advocate|barrister|director|chair|"
         r"president|secretary|fellow|scholar|historian|poet|novelist|artist|activist|contributor|"
         r"correspondent|officer|engineer|doctor|economist|diplomat|consultant|veteran|filmmaker|"
         r"photographer|translator|lawyer|archivist|founder|wife|husband|son|daughter|cousin|brother|sister")
CLUES = [
    re.compile(r"^\s*(--|—|–)\s*\S"),                                    # signed: -- Name
    re.compile(r"\bThe (writer|author|interviewer|columnist)s? (is|are|was)\b", re.I),
    re.compile(r"^[A-Z][\w.'’\- ,]{2,80}\b(is|was|are|were)\b[^.]{0,120}\b(" + ROLES + r")\b", re.I),
    re.compile(r"\b(translated|compiled|edited|curated|written|interviewed|taken) by\b", re.I),
    re.compile(r"^\s*(Curation|Compiled|Text|Interview|Translation)\s*[:\-–]", re.I),
    re.compile(r"\b(talks? to|spoke to|speaks to|in conversation with|in an (exclusive )?interview)\b", re.I),
    re.compile(r"^\s*By\s+[A-Z]"),
]


def evidence(nid, recs=None):
    """Standfirst, first lines, clue sentences near the top/bottom, decoded e-mails."""
    rec = (recs or load_records())[nid]
    with open(text_path(nid), encoding="utf-8") as fh:
        paras = [p.strip() for p in fh.read().split("\n\n", 1)[-1].split("\n") if p.strip()]
    html = _html(rec)
    soup = BeautifulSoup(html, "lxml") if html else None
    deck = ""
    if soup is not None:
        el = soup.select_one('[class*="field-blocknodenewsfield-sub-headline"]')
        deck = el.get_text(" ", strip=True) if el else ""
    clues = []
    for i, p in enumerate(paras):
        near_edge = i < 3 or i >= len(paras) - 5
        for s in re.split(r"(?<=[.!?])\s+", p):
            if any(c.search(s) for c in CLUES) and (near_edge or re.search(r"writer|author|interview|by\b", s, re.I)):
                clues.append(f"[{i + 1}/{len(paras)}] {s.strip()[:300]}")
    return {"deck": deck, "first": paras[:2], "clues": list(dict.fromkeys(clues)),
            "emails": emails(soup) if soup is not None else []}


# --- applying decisions to an exported workbook -----------------------------

def apply(path_in, path_out):
    """Replace generic bylines in every sheet that has Author and Link columns.

    Only the Author cell changes (its formatting is kept) and an 'Author note'
    column is added at the end; nothing else in the workbook is touched.
    """
    from copy import copy
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    fixes = load_fixes()
    wb = load_workbook(path_in)
    counts = {NAMED: 0, EDITORIAL: 0, NOT_NAMED: 0, "undecided": 0}
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if "Author" not in hdr or "Link" not in hdr:
            continue
        a, l = hdr.index("Author") + 1, hdr.index("Link") + 1
        if "Author note" in hdr:
            n = hdr.index("Author note") + 1
        else:
            n = ws.max_column + 1
            head = ws.cell(1, n, "Author note")
            src = ws.cell(1, n - 1)
            head.font, head.alignment, head.fill = copy(src.font), copy(src.alignment), copy(src.fill)
            ws.column_dimensions[get_column_letter(n)].width = 70
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, a)
            if not is_generic(cell.value):
                continue
            fix = fixes.get(node_id(str(ws.cell(r, l).value or "")))
            if fix is None:
                counts["undecided"] += 1
                continue
            counts[fix["status"]] += 1
            cell.value = fix["author"] if fix["status"] == NAMED else "The Daily Star"
            ws.cell(r, n, note(fix)).alignment = Alignment(wrap_text=True, vertical="top")
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(path_out)
    return counts


def main(argv):
    cmd = argv[1] if len(argv) > 1 else "pending"
    if cmd == "apply":
        print(apply(argv[2], argv[3]))
    elif cmd == "evidence":
        recs = load_records()
        for nid in argv[2:]:
            e = evidence(nid, recs)
            print(f"## {nid} | {recs[nid]['title']}")
            print("   standfirst:", e["deck"])
            for p in e["first"]:
                print("   first:", p[:200])
            for c in e["clues"]:
                print("   clue:", c)
            if e["emails"]:
                print("   e-mail:", ", ".join(e["emails"]))
    elif cmd == "pending":
        from .build_xlsx import rows
        fixes = load_fixes()
        items, dupes = rows(apply_fixes=False)
        # Duplicates need a decision too, or a named original stops matching
        # its generic-byline copy and both get listed.
        todo = [i for i in items + dupes if is_generic(i["author"]) and i["nid"] not in fixes]
        for i in todo:
            print(i["nid"], i["date"], i["title"][:80])
        print(f"{len(todo)} generic-byline rows without a decision")


if __name__ == "__main__":
    main(sys.argv)
