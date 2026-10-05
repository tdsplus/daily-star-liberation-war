"""Recall check: are the known-good articles on the Slow Reads - Liberation War sheet?

For each missing one, report the stage where it was lost so the failing
discovery route can be fixed.

Usage:  python -m dstar.recall
"""
import sys

from openpyxl import load_workbook

from .build_xlsx import OUT
from .classify import load_labels
from .extract import load_records
from .store import Candidates, node_id

KNOWN = [
    "https://www.thedailystar.net/slow-reads/focus/news/the-enduring-legacy-the-2nd-bangladesh-liberation-war-course-4241891",
    "https://www.thedailystar.net/slow-reads/focus/news/untold-stories-1971-through-private-letters-4266326",
    "https://www.thedailystar.net/views/in-focus/news/bangladesh-liberation-war-memories-untold-story-2950696",
    "https://www.thedailystar.net/slow-reads/focus/news/birangona-women-bangladesh-3857816",
    "https://www.thedailystar.net/slow-reads/focus/news/the-journey-liberation-war-fifteen-year-old-boy-1506442",
    "https://www.thedailystar.net/slow-reads/focus/news/women-photographers-the-bangladesh-liberation-war-3197181",
    "https://www.thedailystar.net/slow-reads/focus/news/how-does-pakistan-write-1971-4130621",
]


def main():
    ws = load_workbook(OUT)["Slow Reads - Liberation War"]
    in_sheet = {node_id(r[4]) for r in ws.iter_rows(min_row=2, values_only=True) if r[4]}
    cands, recs, labels = Candidates(), load_records(), load_labels()
    missing = 0
    for url in KNOWN:
        nid = node_id(url)
        if nid in in_sheet:
            src = cands.rows.get(nid, {}).get("source_of_discovery", "?")
            print(f"OK       {nid}  (found via {src})")
            continue
        missing += 1
        if nid not in cands.rows:
            why = "never discovered -> discovery routes missed it"
        elif nid not in recs:
            why = "discovered but not extracted (fetch failed or not yet run)"
        elif nid not in labels:
            why = "extracted but not yet classified"
        elif labels[nid]["label"] == "Irrelevant":
            why = f"classified Irrelevant: {labels[nid]['reason']}"
        else:
            why = f"labelled {labels[nid]['label']} but on another sheet ({recs[nid]['canonical_url']})"
        print(f"MISSING  {nid}  {why}\n         {url}")
    print(f"\n{len(KNOWN) - missing}/{len(KNOWN)} known articles present")
    sys.exit(1 if missing else 0)


if __name__ == "__main__":
    main()
